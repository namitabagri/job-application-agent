"""Resume reading (PDF/DOCX/TXT) and resume export (Markdown -> ATS-friendly DOCX)."""
from __future__ import annotations

import io
import re

import docx
import pypdf
from docx.shared import Inches, Pt, RGBColor


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a resume given its filename and raw bytes."""
    name = (filename or "").lower()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    raise ValueError(f"Unsupported resume format: {filename!r} (use PDF, DOCX or TXT)")


def _add_inline(paragraph, text: str) -> None:
    """Render **bold** segments as bold runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = bool(i % 2)


def md_to_docx(md_text: str) -> bytes:
    """Convert simple Markdown to a clean single-column DOCX (ATS-safe: real text,
    standard headings, bullet list style, black Calibri, no tables or images)."""
    document = docx.Document()
    for section in document.sections:
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)

    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            h = document.add_heading(line.lstrip("#").strip(), level=level)
            for run in h.runs:
                run.font.name, run.font.color.rgb = "Calibri", RGBColor(0, 0, 0)
                run.font.size = Pt({1: 18, 2: 12, 3: 11}[level])
        elif line.lstrip().startswith(("- ", "* ")):
            _add_inline(document.add_paragraph(style="List Bullet"), line.lstrip()[2:].strip())
        else:
            _add_inline(document.add_paragraph(), line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
